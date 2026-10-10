"""Prepare the illustrated AROHA report for appending the detailed approach document."""

from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "output" / "report" / "AROHA_Project_Report.docx"
STAGED_REPORT = ROOT / "output" / "report" / "AROHA_Project_Report_with_appendix_divider.docx"


def main() -> None:
    STAGED_REPORT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPORT, STAGED_REPORT)

    document = Document(STAGED_REPORT)
    document.add_page_break()
    title = document.add_paragraph("Appendix A  Professional Approach and Analytics Document")
    title.style = "Heading 1"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "Detailed problem framing, data preparation plan, analytical methodology, validation approach, and responsible-use guidance."
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    document.save(STAGED_REPORT)


if __name__ == "__main__":
    main()
