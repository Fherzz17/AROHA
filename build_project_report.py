from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "report"
OUT.mkdir(parents=True, exist_ok=True)

BLACK = "151515"
INK = "1B1B1B"
WARM = "F2EFE8"
ORANGE = "EE6744"
MUTED = "6F6C66"
LINE = "D9D6CE"
PBLACK = "#" + BLACK
PINK = "#" + INK
PWARM = "#" + WARM
PORANGE = "#" + ORANGE
PMUTED = "#" + MUTED
PLINE = "#" + LINE

def font(size, bold=False, mono=False):
    paths = [
        "C:/Windows/Fonts/consolab.ttf" if bold and mono else "C:/Windows/Fonts/consola.ttf" if mono else "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def rounded(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)

def make_usp_diagram(path):
    im = Image.new("RGB", (1800, 880), PWARM)
    d = ImageDraw.Draw(im)
    title = font(52, bold=True)
    body = font(25)
    small = font(19, bold=True)
    d.text((80, 58), "AROHA decision pipeline", font=title, fill=PINK)
    d.text((82, 128), "The product combines market evidence, transparent scoring and personal profile analysis.", font=body, fill=PMUTED)
    steps = [
        ("01", "Resume and profile", "PDF, DOCX or TXT\nskills extracted locally"),
        ("02", "Market evidence", "15,841 analytics jobs\n1,602 data science jobs"),
        ("03", "AROHA score", "Demand, role coverage,\nsalary and emerging signals"),
        ("04", "Actionable roadmap", "Role fit, skill gaps\nand learning priorities"),
    ]
    x_positions = [80, 510, 940, 1370]
    for idx, (number, heading, detail) in enumerate(steps):
        x = x_positions[idx]
        rounded(d, (x, 280, x + 350, 630), 18, PBLACK if idx in (0, 3) else "#FFFFFF", PORANGE if idx == 2 else PLINE, 3 if idx == 2 else 1)
        d.text((x + 30, 315), number, font=font(20, bold=True), fill=PORANGE)
        d.text((x + 30, 385), heading, font=font(30, bold=True), fill=PWARM if idx in (0, 3) else PINK)
        y = 465
        for line in detail.split("\n"):
            d.text((x + 30, y), line, font=font(20), fill="#D0CCC2" if idx in (0, 3) else PMUTED)
            y += 31
        if idx < len(steps) - 1:
            start = x + 355
            d.line((start, 455, start + 62, 455), fill=PORANGE, width=4)
            d.polygon([(start + 62, 455), (start + 48, 446), (start + 48, 464)], fill=PORANGE)
    d.text((82, 740), "USP: AROHA does not stop at market dashboards. It converts evidence into a concrete next skill decision for an individual user.", font=font(25), fill=PINK)
    im.save(path)

def make_code_image(source_path, title, start_marker, end_marker, out_path):
    source = source_path.read_text(encoding="utf-8")
    snippet = source[source.index(start_marker):]
    if end_marker and end_marker in snippet:
        snippet = snippet[:snippet.index(end_marker)]
    lines = snippet.strip().splitlines()[:34]
    im = Image.new("RGB", (1800, max(520, 150 + len(lines) * 31)), "#101010")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1800, 90), fill="#191919")
    d.ellipse((44, 32, 58, 46), fill="#EE6744")
    d.ellipse((68, 32, 82, 46), fill="#E8B64B")
    d.ellipse((92, 32, 106, 46), fill="#77BE70")
    d.text((140, 26), title, font=font(25, bold=True), fill="#F2EFE8")
    for i, line in enumerate(lines):
        y = 120 + i * 31
        d.text((42, y), f"{i + 1:02}", font=font(17, mono=True), fill="#6F6C66")
        color = "#F2EFE8"
        stripped = line.lstrip()
        if stripped.startswith(("def ", "async def ", "return ", "if ", "elif ", "else", "from ", "import ")):
            color = "#FF9D7D"
        elif "#" in stripped:
            color = "#8A9584"
        elif any(token in stripped for token in ("\"", "'")):
            color = "#CBE895"
        d.text((105, y), line.replace("\t", "    "), font=font(17, mono=True), fill=color)
    im.save(out_path)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Arial"
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(18 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    run = p.add_run(text)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    return p

def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.22
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.size = Pt(10.5)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor.from_string(INK)
    return p

def add_bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(value)
        run.font.size = Pt(10.5)
        run.font.name = "Arial"

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(MUTED)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i], BLACK)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=9)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            shade(cells[i], "F7F4ED" if row_index % 2 == 0 else "FFFFFF")
            set_cell_text(cells[i], str(value), size=8.9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def main():
    usp = OUT / "aroha-usp-architecture.png"
    score_code = OUT / "aroha-score-code.png"
    resume_code = OUT / "aroha-resume-code.png"
    client_code = OUT / "aroha-client-controller-code.png"
    recommendation_code = OUT / "aroha-recommendation-code.png"
    make_usp_diagram(usp)
    make_code_image(ROOT / "services" / "scoring.py", "services/scoring.py - explainable score", "def calculate_skillpulse_score", "def re_escape", score_code)
    make_code_image(ROOT / "backend" / "app.py", "backend/app.py - resume skill extraction", "@app.post(\"/api/extract-resume\", response_model=ExtractResponse)", "@app.get(\"/api/dataset/sample\")", resume_code)
    make_code_image(ROOT / "frontend" / "static" / "js" / "app.js", "frontend/static/js/app.js - resume upload workflow", "function setupResumeUpload", "function renderSelectedUserSkillsChips", client_code)
    make_code_image(ROOT / "services" / "recommender_service.py", "services/recommender_service.py - recommendation logic", "def recommend_skills", None, recommendation_code)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.72)
    sec.right_margin = Inches(.72)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Title"].font.size = Pt(33)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True

    p = doc.add_paragraph(style="Title")
    p.add_run("AROHA Project Report")
    p.paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph()
    p.add_run("AI powered skill intelligence and career decision support").font.size = Pt(15)
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    p.paragraph_format.space_after = Pt(20)
    add_body(doc, "Team Syntax Society | SAS Hackathon 2026 | AI and Workforce Analytics")
    add_body(doc, "AROHA is a decision-support MVP that turns job-market evidence and an individual profile into a focused recommendation for what to learn next.")
    doc.add_picture(str(ROOT / "mvp-home.png"), width=Inches(6.95))
    add_caption(doc, "Figure 1. AROHA MVP home screen and career-signal interface.")

    add_heading(doc, "Project Summary")
    add_body(doc, "AROHA addresses a practical career-planning problem: candidates often see job listings and skill lists but cannot determine which capability to prioritise. The MVP combines market evidence, transparent scoring, machine-learning outputs and an individual profile into a single decision flow.")
    add_table(doc, ["Area", "MVP implementation", "User value"], [
        ("Skill intelligence", "Normalises and analyses job-posting skills", "Clarifies what employers are asking for"),
        ("Role fit", "Compares existing skills with target-role requirements", "Makes gaps visible and prioritised"),
        ("Resume import", "Reads PDF, DOCX or TXT in memory and extracts indexed skills", "Removes manual profile entry"),
        ("Explainable score", "Weights demand, role coverage, salary association and frontier signals", "Shows why a skill is recommended"),
        ("Learning roadmap", "Ranks next skills with evidence", "Converts analysis into action"),
    ], [1.25, 3.0, 2.45])

    add_heading(doc, "Problem and Target User")
    add_body(doc, "Students, early-career professionals and placement stakeholders need a credible way to translate fragmented job-market information into a personal plan. Existing job boards show vacancies; AROHA explains the capability signals behind those vacancies and maps them to a target role.")
    add_bullets(doc, [
        "Candidate: identify the few skills with the strongest potential impact for a target role.",
        "Placement or educator: understand recurring employer requirements and curriculum gaps.",
        "Recruiter or workforce planner: inspect market demand, role patterns and salary context.",
    ])

    add_heading(doc, "MVP Demonstration")
    add_body(doc, "The implemented MVP is a local web application with a navigable product interface, interactive role and skill exploration, resume upload, recommendation analysis and API-backed data services.")
    add_table(doc, ["Workflow", "What happens", "Evidence in MVP"], [
        ("1. Import", "User uploads a resume or enters skills", "Resume upload accepts PDF, DOCX and TXT files up to 5 MB"),
        ("2. Extract", "Text is parsed in memory and passed to the existing skill extractor", "Detected skills are added to the user profile"),
        ("3. Compare", "Profile is compared against a selected career role", "Match score, missing core skills and radar chart update"),
        ("4. Act", "Recommendation service ranks next capabilities", "Learning roadmap explains what to learn and why"),
    ], [1.0, 3.25, 2.45])

    add_heading(doc, "Unique Selling Proposition")
    add_body(doc, "The core USP is the bridge from market data to an individual next action. AROHA combines observable hiring evidence with an explainable score and then applies that evidence to a specific user profile. This is more useful than a static demand dashboard and more credible than an unexplained generic recommendation.")
    doc.add_picture(str(usp), width=Inches(6.95))
    add_caption(doc, "Figure 2. AROHA USP: evidence flows from resume and market data to a specific learning decision.")
    add_bullets(doc, [
        "Explainable: the score exposes its four component signals and weights.",
        "Personal: resume extraction removes manual setup and starts from the candidate's actual profile.",
        "Evidence-based: recommendations are tied to demand, role coverage and salary context.",
        "Responsible: frontier signals are labelled as analytical indicators rather than fabricated historical facts.",
    ])

    add_heading(doc, "Data and Analytical Method")
    add_table(doc, ["Dataset", "Records", "Use in AROHA"], [
        ("Analytics Jobs", "15,841", "Skill frequency, role coverage, location and salary parsing"),
        ("DataScience Jobs", "1,602", "Employer, role and compensation intelligence"),
        ("JDS Skill Traits", "139", "Junior salary-hike classification and feature analysis"),
        ("SDS Personality Traits", "161", "Senior success-classification exploratory model"),
    ], [2.2, 1.0, 3.5])
    add_body(doc, "The AROHA score uses a transparent weighted formula: 35 percent demand, 25 percent role coverage, 25 percent salary association and 15 percent emerging-skill alignment. Salary signals are described as associations, not causal effects.")

    add_heading(doc, "Implementation Evidence")
    add_heading(doc, "Explainable Skill Score", level=2)
    add_body(doc, "The score service explicitly derives its components, applies the documented weights and returns both the score and the underlying empirical metrics. This makes the product result inspectable during a demo.")
    doc.add_picture(str(score_code), width=Inches(6.95))
    add_caption(doc, "Figure 3. Production code excerpt for the transparent weighted skill score.")
    add_heading(doc, "Resume Import and Extraction", level=2)
    add_body(doc, "The API validates file type and size, reads the document in memory, extracts text with format-specific readers, and routes the text through the existing normalisation pipeline. The endpoint does not persist uploaded resumes.")
    doc.add_picture(str(resume_code), width=Inches(6.95))
    add_caption(doc, "Figure 4. Production code excerpt for secure in-memory resume skill extraction.")

    add_heading(doc, "Technical Architecture")
    add_table(doc, ["Layer", "Technology", "Responsibility"], [
        ("Frontend", "HTML, CSS and JavaScript", "Interactive AROHA interface, charts, profile flow and resume upload"),
        ("API", "FastAPI and Pydantic", "Typed endpoints for analytics, extraction and recommendations"),
        ("Analytics", "Pandas and scikit-learn", "Data preparation, market intelligence and predictive models"),
        ("Storage", "SQLite", "Local job-market data access"),
        ("Document parsing", "pypdf and python-docx", "In-memory PDF, DOCX and TXT skill extraction"),
    ], [1.2, 2.0, 3.5])

    add_heading(doc, "Evaluation and Responsible Use")
    add_body(doc, "The project includes automated tests for extraction, scoring, analytics, recommendation services and FastAPI endpoints. The documented ML evaluation compares Logistic Regression and Random Forest models, with the Junior Data Scientist model reporting 88.6 percent accuracy and 0.889 F1 for the stated validation split.")
    add_bullets(doc, [
        "Forecast and frontier signals are shown as indicators, not guarantees.",
        "Skill and salary relationships are communicated as observational associations.",
        "Resume files are processed in memory for extraction and are not stored by the endpoint.",
        "Personality outputs are labelled with ethical limits and are not presented as psychological diagnosis.",
    ])

    add_heading(doc, "Conclusion")
    add_body(doc, "AROHA is a complete MVP rather than a static concept: it accepts a user profile, reads a resume, analyses market evidence, identifies role gaps and produces a concrete learning roadmap. Its distinctive strength is the transparent connection between data, model outputs and a user decision.")

    doc.add_page_break()
    add_heading(doc, "Feature Catalogue and Screen Evidence")
    add_body(doc, "This section documents the complete user-facing MVP. Screens are captured from the running AROHA application at the local product viewport. Each feature is connected to the implementation layer responsible for it.")
    features = [
        ("Home", "01-home.png", "Sets the product context and gives users direct paths to profile analysis or market exploration.", "frontend/templates/index.html and frontend/static/css/style.css"),
        ("Market Dashboard", "02-dashboard.png", "Presents current job-market signals, in-demand skills and industry spread using API-backed charts.", "frontend/static/js/app.js and backend/analytics_service.py"),
        ("Skill Search", "03-skill-search.png", "Lets users inspect demand score, observed job presence, trends, related skills and top industries for a selected technology.", "frontend/static/js/app.js and backend/analytics_service.py"),
        ("Career Explorer", "04-career-explorer.png", "Compares job roles through salary, skills, locations and role-level evidence.", "frontend/static/js/app.js and backend/analytics_service.py"),
        ("Emerging Skills", "05-emerging-skills.png", "Shows frontier capabilities with adoption labels and explicit forecast caveats.", "backend/analytics_service.py"),
        ("Skill Gap and Resume Import", "06-skill-gap-resume.png", "Accepts a candidate profile or resume, identifies skills, calculates alignment and returns a structured roadmap.", "backend/app.py, services/recommender_service.py and frontend/static/js/app.js"),
        ("NLP Job Analyzer", "07-nlp-analyzer.png", "Extracts and normalises skills from a pasted job description, then suggests matching roles.", "services/skill_extractor.py and backend/app.py"),
        ("About and Methodology", "08-methodology.png", "Explains source data, analytical scope and the project methodology for judges and stakeholders.", "frontend/templates/index.html and services/data_service.py"),
    ]
    for index, (name, image, explanation, code_ref) in enumerate(features, 1):
        add_heading(doc, f"{index}. {name}", level=2)
        add_body(doc, explanation)
        add_body(doc, f"Implementation: {code_ref}", bold_lead="Implementation:")
        doc.add_picture(str(OUT / "screenshots" / image), width=Inches(6.75))
        add_caption(doc, f"Figure {4 + index}. AROHA {name} feature screen.")

    doc.add_page_break()
    add_heading(doc, "Feature to Code Traceability")
    add_body(doc, "The following map makes it straightforward for judges to connect the visible product capabilities with the relevant source code.")
    add_table(doc, ["Feature", "Primary source files", "Key implementation"], [
        ("Navigation and views", "frontend/templates/index.html; frontend/static/js/app.js", "switchView controls the single-page application views"),
        ("Dashboard and charts", "frontend/static/js/app.js; frontend/static/js/charts.js; backend/analytics_service.py", "Fetches API summaries and renders interactive charts"),
        ("Skill intelligence", "services/scoring.py; backend/analytics_service.py", "Computes explainable score and market metrics"),
        ("Role analysis", "backend/analytics_service.py; utils/constants.py", "Returns role requirements, salary and skill evidence"),
        ("Resume upload", "backend/app.py; frontend/static/js/app.js", "Validates document, extracts text in memory and maps skills"),
        ("Recommendations", "services/recommender_service.py", "Ranks capability gaps and creates learning roadmap"),
        ("NLP extraction", "services/skill_extractor.py; utils/normalizer.py", "Extracts aliases and normalises canonical skills"),
    ], [1.25, 2.85, 2.4])
    add_heading(doc, "Client Workflow Code", level=2)
    add_body(doc, "The browser-side upload workflow validates the selected file, posts it as multipart form data and refreshes the user profile with detected skills before requesting a new recommendation.")
    doc.add_picture(str(client_code), width=Inches(6.75))
    add_caption(doc, "Figure 13. Client-side resume upload and profile refresh workflow.")
    add_heading(doc, "Recommendation Service Code", level=2)
    add_body(doc, "The recommendation service compares the profile with the selected target role and returns prioritised gaps together with evidence-oriented learning steps.")
    doc.add_picture(str(recommendation_code), width=Inches(6.75))
    add_caption(doc, "Figure 14. Server-side recommendation service excerpt.")

    doc.add_page_break()
    add_heading(doc, "User Manual")
    add_body(doc, "This manual describes how a judge or first-time user can use the MVP from the home screen to a personalised recommendation.")
    add_heading(doc, "Starting the Application", level=2)
    add_bullets(doc, [
        "Install requirements with: python -m pip install -r requirements.txt.",
        "Run the FastAPI experience with: python run.py --fastapi.",
        "Open http://127.0.0.1:8000 in a browser.",
        "Use the Home page to begin with either Analyse my profile or Explore the market.",
    ])
    add_heading(doc, "Using Market Intelligence", level=2)
    add_bullets(doc, [
        "Open Dashboard to review high-level market volume, leading skills and industry demand charts.",
        "Open Skill Search, enter a skill such as Python or SQL, then inspect its demand, related skills and forecast panel.",
        "Open Career Explorer, select a target role and review salary, core skills and location information.",
        "Open Emerging Skills to review frontier technologies. Treat these as labelled analytical signals rather than guarantees.",
    ])
    add_heading(doc, "Creating a Personal Skill Gap Analysis", level=2)
    add_bullets(doc, [
        "Open Recommendations from the left navigation.",
        "Choose the target career role from the first selector.",
        "Upload a PDF, DOCX or TXT resume up to 5 MB, or add skills manually.",
        "After extraction, confirm the skill chips that were added to your profile. Remove a chip if it does not apply.",
        "Select Run Skill Gap Analysis.",
        "Review role alignment, missing core capabilities, frontier skills, the benchmark radar and the learning roadmap.",
    ])
    add_heading(doc, "Analysing a Job Description", level=2)
    add_bullets(doc, [
        "Open NLP Job Analyzer.",
        "Paste a job description into the text area and select Extract and Normalize Skills.",
        "Use the returned canonical skills and matching-role suggestions to compare the listing with your current profile.",
    ])
    add_heading(doc, "Data and Privacy Notes", level=2)
    add_bullets(doc, [
        "Uploaded resumes are read in memory for extraction and are not stored by the API endpoint.",
        "Text-based PDFs work best. Scanned image-only PDFs may not provide readable text.",
        "The tool provides decision support, not a hiring or psychological assessment. Salary relationships are observational, not causal.",
    ])

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AROHA Project Report | Syntax Society | SAS Hackathon 2026").font.size = Pt(8)
    output = OUT / "AROHA_Project_Report.docx"
    doc.save(output)
    print(output)

if __name__ == "__main__":
    main()
