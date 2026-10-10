"""Tighten the final glossary table so it remains a single readable page."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


PATH = Path(__file__).resolve().parent / "output" / "report" / "AROHA_Combined_Hackathon_Submission.docx"


def main() -> None:
    document = Document(PATH)
    glossary = document.tables[-1]
    for row in glossary.rows:
        tr_pr = row._tr.get_or_add_trPr()
        for element in list(tr_pr):
            if element.tag == qn("w:cantSplit"):
                tr_pr.remove(element)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    document.save(PATH)


if __name__ == "__main__":
    main()
